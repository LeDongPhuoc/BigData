"""Create a concise, editable 20-slide presentation inside the HUIT template."""
from pathlib import Path
import shutil
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/slides-assets'
OUT.mkdir(exist_ok=True)
TARGET=Path(r'C:\Users\NGUYEN THANH NAM\Downloads\vie_16x9_slide-nhan-dien-huit.pptx')
BACKUP=TARGET.with_name(TARGET.stem+'_Mau_goc.pptx')
if not BACKUP.exists():shutil.copy2(TARGET,BACKUP)
p=Presentation(BACKUP)
for sid in list(p.slides._sldIdLst):
    p.part.drop_rel(sid.rId)
    p.slides._sldIdLst.remove(sid)
BLUE='005BAA';NAVY='123653';LIGHT='EDF4FA';GRAY='526575';TEAL='14695D';GOLD='C08A00';RED='EA1D25'
manifest=[]

def rgb(v):return RGBColor.from_string(v)

def box(s,x,y,w,h,text='',size=24,color=NAVY,bold=False,fill=None,align=PP_ALIGN.LEFT):
    if fill:
        sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background()
    else:sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=Inches(.12 if fill else 0);tf.margin_right=Inches(.12 if fill else 0)
    tf.margin_top=Inches(.08 if fill else 0);tf.margin_bottom=Inches(.04)
    for i,line in enumerate(text.split('\n')):
        pa=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        pa.alignment=align;pa.space_after=Pt(6);pa.space_before=Pt(0)
        pa.line_spacing=1.12
        r=pa.add_run();r.text=line;r.font.name='Arial';r.font.size=Pt(size)
        r.font.bold=bold;r.font.color.rgb=rgb(color)
    return sh

def slide(title,notes='',cover=False):
    s=p.slides.add_slide(p.slide_layouts[0 if cover else 5])
    for sh in list(s.shapes):
        el=sh._element;el.getparent().remove(el)
    if not cover:
        box(s,.55,1.03,12.05,.72,title,30,BLUE,True)
        box(s,11.65,.18,1.05,.32,f'{len(p.slides):02d} / 20',14,'FFFFFF',align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes
    manifest.append({'slide':len(p.slides),'title':title,'speaker_notes':notes})
    return s

def note(s,content):box(s,.6,6.58,12,.38,content,15,GRAY)

def bullets(s,items,x=.7,y=2.1,w=11.8,size=27,gap=1.05):
    for i,t in enumerate(items):
        box(s,x,y+i*gap,.25,.65,'•',size,BLUE,True)
        box(s,x+.4,y+i*gap,w-.4,gap-.1,t,size)

def cards(s,items,y=2.15):
    w=3.88
    for i,(title,body) in enumerate(items):
        x=.6+i*4.12
        box(s,x,y,w,3.7,fill=LIGHT)
        box(s,x+.2,y+.3,w-.4,.95,title,27,BLUE,True)
        box(s,x+.2,y+1.4,w-.4,1.95,body,23)

def node(s,x,y,w,h,label,color=BLUE,size=23):
    sh=box(s,x,y,w,h,label,size,color,True,LIGHT,PP_ALIGN.CENTER)
    sh.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    return sh

def arrow(s,x1,y1,x2,y2,label=None,lx=None,ly=None,lw=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(BLUE);sh.line.width=Pt(2)
    end=OxmlElement('a:tailEnd');end.set('type','triangle');sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(end)
    if label:box(s,lx if lx is not None else (x1+x2)/2-.7,ly if ly is not None else (y1+y2)/2-.4,lw,.45,label,18,GRAY,align=PP_ALIGN.CENTER)

def pic(s,path,x,y,w,h):
    from PIL import Image
    im=Image.open(path);iw,ih=im.size
    ratio=min(w/iw,h/ih);rw,rh=iw*ratio,ih*ratio
    s.shapes.add_picture(str(path),Inches(x+(w-rw)/2),Inches(y+(h-rh)/2),width=Inches(rw),height=Inches(rh))

s=slide('Ứng dụng ArangoDB',cover=True,notes='Giới thiệu tên đề tài và sản phẩm Job Referral Network. Điền thông tin nhóm, MSSV, lớp và giảng viên trước buổi trình bày; các thông tin này chưa được cung cấp. Bài trình bày dự kiến 12–15 phút, dành thêm thời gian nếu demo trực tiếp. Ngày chốt nội dung 12/10/2026.')
box(s,.75,1.1,11.85,.84,'ỨNG DỤNG ARANGODB',37,RED,True,align=PP_ALIGN.CENTER)
box(s,.9,1.93,11.55,1.25,'GỢI Ý VIỆC LÀM\nVÀ ĐƯỜNG KẾT NỐI GIỚI THIỆU',32,BLUE,True,align=PP_ALIGN.CENTER)
box(s,1.2,3.32,10.95,.48,'Job Referral Network  •  Khoa Công nghệ Thông tin',23,align=PP_ALIGN.CENTER)
box(s,1.2,4.05,10.95,.6,'Nhóm / MSSV / Lớp / GVHD: [Cần bổ sung]',20,GRAY,align=PP_ALIGN.CENTER)
box(s,1.2,4.73,10.95,.36,'TP. Hồ Chí Minh, tháng 10 năm 2026',18,GRAY,align=PP_ALIGN.CENTER)

s=slide('BÀI TOÁN VÀ MỤC TIÊU','Mô hình nghiệp vụ giả định, chưa có khảo sát phỏng vấn doanh nghiệp cụ thể. Hai bài toán riêng: đối chiếu hồ sơ với công việc và tìm người giới thiệu qua mạng lưới. Source: detai.md; docs/requirements.md; báo cáo chương 1–2.')
cards(s,[('Việc nào phù hợp?','Đối chiếu kỹ năng,\nkinh nghiệm và mong muốn.'),('Có thể nhờ ai?','Tìm người giới thiệu\nqua các kết nối hợp lệ.'),('Vì sao được gợi ý?','Hiển thị điểm thành phần\nvà đường đi có giải thích.')])
note(s,'Mục tiêu: ứng dụng chạy thật với giao diện → API → ArangoDB.')

s=slide('ARANGODB LÀ GÌ?','ArangoDB 3.12 lưu các thực thể như JSON documents và quan hệ qua edges trong property graph. Cạnh _from/_to luôn có chiều nhưng traversal ANY có thể đi hai chiều. Named graph khai báo loại quan hệ. Nguồn chính thức: https://docs.arango.ai/arangodb/3.12/graphs/ ; https://docs.arango.ai/arangodb/3.12/aql/ . Truy cập khi biên soạn báo cáo ngày 12/10/2026.')
cards(s,[('Document','Lưu người dùng, công ty,\nviệc làm và kỹ năng.'),('Graph','Lưu quan hệ bằng edge;\nkhai báo named graph.'),('AQL','Lọc dữ liệu, tính điểm\nvà duyệt mạng lưới.')])
note(s,'Phiên bản của dự án: ArangoDB 3.12.12, single server.')

s=slide('VÌ SAO CHỌN ARANGODB?','Lựa chọn dựa trên sự phù hợp mô hình của bài toán, không phải phép so sánh hiệu năng với SQL hoặc Neo4j. Named graph không tự bảo đảm toàn vẹn khi mọi thao tác đều đi qua raw document API: dự án kiểm tra bổ sung tại service/repository. Nguồn: docs/database/graph-schema.md; docs/database/query-design.md; tài liệu Graphs chính thức.')
bullets(s,['Document và quan hệ nằm trong cùng cơ sở dữ liệu.','AQL vừa đối chiếu hồ sơ vừa tìm đường nhiều bước.','Index và transaction hỗ trợ truy vấn, ghi nhất quán.'])
note(s,'Chưa có benchmark đối chứng để kết luận nhanh hơn các DBMS khác.')

s=slide('KIẾN TRÚC HỆ THỐNG','React/TypeScript dùng TanStack Query; Nginx proxy sang FastAPI; backend chia API/schema/service/repository; python-arango gọi AQL/transaction. Frontend không truy cập database trực tiếp. Nguồn: compose.yaml; backend/modules/; frontend/src/; docs/architecture.md.')
for x,title in [(.6,'React\nGiao diện'),(3.85,'Nginx\nProxy'),(7.1,'FastAPI\nNghiệp vụ'),(10.35,'ArangoDB\nAQL / Graph')]:node(s,x,2.9,2.4,1.4,title)
for x in [3,6.25,9.5]:arrow(s,x,3.6,x+.82,3.6)
box(s,.8,5.08,11.8,.8,'FastAPI kiểm tra quyền  •  ArangoDB thực thi truy vấn và transaction',24,align=PP_ALIGN.CENTER)
note(s,'Triển khai local bằng Docker Compose; backend dùng DB runtime user riêng.')

s=slide('MÔ HÌNH GRAPH CỐT LÕI','Named graph: jobReferralGraph. connections chỉ một cạnh/cặp canonical người dùng, traversal ANY. works_at: user→company; has_skill: user→skill; job_requires_skill: job→skill; job_at_company: job→company. Không suy ra quen biết từ cùng công ty/kỹ năng. Nguồn: database/migrations/; docs/database/graph-schema.md.')
node(s,.65,2.25,2.7,1,'users\nỨng viên')
node(s,7.75,2.25,2.9,1,'users\nNgười giới thiệu')
node(s,.65,4.6,2.7,1,'skills\nKỹ năng')
node(s,5.2,4.6,2.25,1,'jobs\nViệc làm')
node(s,9.55,4.6,2.9,1,'companies\nDoanh nghiệp')
arrow(s,3.35,2.75,7.75,2.75,'connections · ANY',lx=4,ly=2.15,lw=3.2)
arrow(s,2,3.25,2,4.6,'has_skill',lx=2.15,ly=3.72,lw=1.8)
arrow(s,9.2,3.25,10.6,4.6,'works_at',lx=10.2,ly=3.7,lw=1.7)
arrow(s,5.2,5.1,3.35,5.1)
box(s,3.4,5.55,3,.56,'job_requires_skill',16,GRAY)
arrow(s,7.45,5.1,9.55,5.1)
box(s,7.55,5.55,2.3,.56,'job_at_company',16,GRAY)
note(s,'Đỉnh lưu thực thể; cạnh lưu quan hệ và thuộc tính nghiệp vụ.')

s=slide('DỮ LIỆU: CÔNG KHAI VÀ KIỂM THỬ','Snapshot ngày 12/10/2026: KMS 5 jobs, Bosch 12, Grab 12; 28 skills. Dữ liệu chuẩn hóa không đại diện toàn bộ tin tuyển dụng của các công ty. Tin thiếu lương/kinh nghiệm không được điền giả; skills trích từ nguồn lưu preferred nếu chưa xác định required. Công ty public chưa có nhân viên được xác minh trong ứng dụng. Nguồn: docs/public-data.md; database/catalog/public_jobs_2026-10-12.json.')
cards(s,[('29 tin công khai','KMS 5  •  Bosch 12  •  Grab 12\nCó URL và thời điểm lấy tin.'),('28 kỹ năng','Chuẩn hóa danh mục\nđể đối chiếu hồ sơ.'),('Graph kiểm thử','Hồ sơ / employment giả lập\nđược gắn nhãn riêng.')])
note(s,'Snapshot ngày 12/10/2026; không tự tạo nhân viên giả cho công ty thật.')

s=slide('GỢI Ý VIỆC LÀM: ĐIỂM CÓ GIẢI THÍCH','Công thức database/queries.py và database/aql/matching_rows.aql: S=.4K+.2E+.15L+.15V+.1W. K dùng required/preferred importance, xét proficiency/years; nếu có cả hai: .8R+.2P. E=min(1,years/min_years), min=0→1, dữ liệu yêu cầu thiếu→0. Remote hoặc location bằng nhau→L=1. Level exact→V=1. Work mode trong preference→W=1. Salary_overlap chỉ diagnostic. Đây là quy tắc, chưa được hiệu chỉnh theo labels tuyển dụng. Eligible khác matched_only.')
box(s,.8,2.03,11.8,.7,'Điểm = 40% K + 20% E + 15% L + 15% V + 10% W',29,BLUE,True,align=PP_ALIGN.CENTER)
labels=[('Kỹ năng','40%'),('Kinh nghiệm','20%'),('Địa điểm','15%'),('Cấp độ','15%'),('Hình thức','10%')]
for i,(label,value) in enumerate(labels):
    x=.7+i*2.55
    box(s,x,3.3,2.28,1.6,fill=LIGHT)
    box(s,x+.1,3.52,2.08,.65,value,35,BLUE,True,align=PP_ALIGN.CENTER)
    box(s,x+.08,4.3,2.12,.4,label,19,align=PP_ALIGN.CENTER)
box(s,.9,5.32,11.5,.74,'Khớp kỹ năng theo mức thành thạo và số năm kinh nghiệm.',24,align=PP_ALIGN.CENTER)
note(s,'Điểm xếp hạng theo dữ liệu đã lưu; không phải xác suất được tuyển.')

s=slide('TÌM ĐƯỜNG TỚI NGƯỜI GIỚI THIỆU','Ví dụ fixture A→B→C với C làm tại X current/verified. Hops chỉ đếm connections, không tính works_at. Tối đa 3 hops; ANY trên named graph nhưng chỉ duyệt connections/users. Endpoint cần employment current/verified và visibility. Khi không có nhân viên public được xác minh, danh sách rỗng là đúng. Nguồn: path_rows.aql; employees_rows.aql; arango-behavior-proof.json.')
for x,label in [(.7,'A\nỨng viên'),(3.65,'B\nKết nối'),(6.6,'C\nNgười giới thiệu'),(10,'X\nDoanh nghiệp')]:node(s,x,2.85,2.5,1.4,label)
arrow(s,3.2,3.55,3.62,3.55)
arrow(s,6.15,3.55,6.57,3.55)
arrow(s,9.1,3.55,9.97,3.55)
box(s,2.1,4.65,5.5,.55,'A → B → C: 2 kết nối',29,BLUE,True,align=PP_ALIGN.CENTER)
box(s,9.2,4.43,3.6,1.1,'works_at\nHiện tại + đã xác minh',20,GRAY,align=PP_ALIGN.CENTER)
note(s,'Chỉ đi qua kết nối accepted, tài khoản active và quyền chia sẻ hợp lệ.')

s=slide('AQL: DUYỆT GRAPH VÀ LOẠI NHÁNH','Đoạn code là trích rút để giải thích, không phải query đầy đủ. PRUNE phải có guard e!=null ở root. FILTER !invalid loại node tại điểm dừng. uniqueVertices=path giữ đường thay thế và tránh lặp đỉnh trên đường. Nguồn chính xác: database/aql/path_rows.aql và database/queries.py. Tài liệu: https://docs.arango.ai/arangodb/3.12/aql/graph-queries/traversals/ .')
box(s,.7,2.1,7.4,3.7,fill=LIGHT)
code='FOR v, e, p IN 1..@max_depth\n  ANY @actor_id GRAPH "jobReferralGraph"\n  PRUNE invalid = /* status + privacy */\n  OPTIONS { edgeCollections: ["connections"],\n            uniqueVertices: "path" }\n  FILTER !invalid'
sh=box(s,.92,2.43,6.96,2.95,code,20)
for pa in sh.text_frame.paragraphs:
    for r in pa.runs:r.font.name='Consolas'
bullets(s,['Độ sâu ≤ 3','PRUNE nhánh sai','Giữ đường thay thế'],x=8.55,y=2.52,w=4.1,size=23,gap=1)
note(s,'Đoạn AQL rút gọn; query đầy đủ còn kiểm tra block, privacy và employment.')

s=slide('XẾP HẠNG ĐƯỜNG GIỚI THIỆU','D=.6×mean declared trust+.4×1/hops. Fixture 2 hops mean trust=.85→.71. Đường 3 hops có .673333. Missing trust=.5; không phải xác suất. Shortest mode sort hops trước score; mặc định score trước hops. Jobs_with_referral giữ eligible rồi sort job_score trước path_score. Nguồn: path_rows.aql; jobs_with_referral.aql; arango-behavior-proof.json.')
box(s,.7,2.08,11.9,.72,'Điểm đường = 60% × trust trung bình + 40% × (1 / hops)',28,BLUE,True,align=PP_ALIGN.CENTER)
node(s,.9,3.55,4.3,1.5,'2 hops  •  trust = 0,85')
arrow(s,5.2,4.3,6.3,4.3)
box(s,6.55,3.38,5.4,1.65,'71 / 100',51,BLUE,True,LIGHT,PP_ALIGN.CENTER)
box(s,1,5.62,11.5,.48,'Cơ hội có đường: ưu tiên đủ điều kiện và điểm việc làm trước.',23,align=PP_ALIGN.CENTER)
note(s,'Trust là độ tin cậy khai báo; điểm đường không phải tỷ lệ giới thiệu thành công.')

s=slide('SẢN PHẨM: GỢI Ý THEO HỒ SƠ','Ảnh artifact browser trên fixture, không phải nhân viên hoặc vị trí tuyển dụng thực đã xác minh. Dashboard gọi /matching/jobs; /jobs?view=recommended hỗ trợ filter/pagination. Thêm/xóa skill invalidates query cache; headline không tự chuyển thành skill. matched_only=true yêu cầu ít nhất một skill matched. Nguồn: docs/reports/frontend-dashboard-1440.png; recommendations-fix-2026-10-12.md.')
from PIL import Image
Image.open(ROOT/'docs/reports/frontend-dashboard-1440.png').crop((265,675,975,1055)).save(OUT/'recommendation-detail.png')
pic(s,OUT/'recommendation-detail.png',.55,1.92,8.2,4.53)
bullets(s,['Điểm phù hợp','Kỹ năng khớp','Bộ lọc và phân trang'],x=9.02,y=2.37,w=3.6,size=23,gap=1.1)
note(s,'Ảnh sản phẩm chạy với dữ liệu kiểm thử; kết quả đọc từ API/ArangoDB thật.')

s=slide('SẢN PHẨM: ĐƯỜNG KẾT NỐI','Ảnh artifact frontend-paths-1440.png trên fixture. React Flow vẽ vertices và edges từ response /graph/referral-paths; không tự tính đường tại browser. Hiển thị hai đường thay thế, hops, trust, điểm và current_verified_employee. Referrer C trong ảnh là kiểm thử. Source: docs/reports/frontend-paths-1440.png; docs/arango-proof-guide.md.')
pic(s,ROOT/'docs/report-assets/paths.png',.55,1.91,8.2,4.57)
bullets(s,['Graph từ AQL','Số bước + giải thích','Gửi yêu cầu giới thiệu'],x=9.02,y=2.38,w=3.6,size=23,gap=1.1)
note(s,'React Flow hiển thị đường; ArangoDB tìm đường và kiểm tra quan hệ.')

s=slide('QUY TRÌNH GIỚI THIỆU','Đường chính pending→accepted→submitted→completed. pending có thể rejected/cancelled/expired; accepted có thể cancelled/expired. Candidate hủy; referrer xử lý. Create/accept/submit kiểm tra lại policy/job/employment/path. Submitted lưu evidence_ref, chưa gửi email hoặc ATS. Ghi referral/history/notification/audit trong transaction, revision If-Match và idempotency chống trùng. Nguồn: backend/modules/referrals/; docs/backend.md.')
for x,title in [(.6,'pending\nGửi yêu cầu'),(3.85,'accepted\nChấp nhận'),(7.1,'submitted\nXác nhận đã gửi'),(10.35,'completed\nHoàn tất')]:node(s,x,2.7,2.4,1.5,title,size=22)
for x in [3,6.25,9.5]:arrow(s,x,3.45,x+.82,3.45)
box(s,.95,4.75,11.4,.62,'Có nhánh từ chối / hủy / hết hạn theo trạng thái và quyền.',25,align=PP_ALIGN.CENTER)
box(s,.95,5.55,11.4,.5,'Transaction + kiểm tra lại đường + chống gửi trùng',24,BLUE,True,align=PP_ALIGN.CENTER)
note(s,'MVP lưu xác nhận / mã bằng chứng; chưa tích hợp email, tải CV hoặc ATS.')

s=slide('CHỨNG MINH ARANGODB THỰC SỰ THAM GIA','Artifact arango-behavior-proof.json ngày 12/10/2026: baseline 2/3 hops. after_block_b_c còn 3 hops; after_b_disallows_paths còn 3 hops; after_c_employment_unverified=[]; plan có TraversalNode/IndexNode và company_current_employees; schema error1620 transaction_rolled_back=true; main_database_modified=false và isolated_database_cleaned_up=true. Script scripts/prove_arangodb.py dùng database UUID thật, không mutate main.')
rows=[('Chặn cạnh B–C','Đường 2 bước mất; còn đường 3 bước.'),('B tắt quyền chia sẻ','Nhánh qua B bị loại.'),('C chưa xác minh employment','Không còn đường hợp lệ.'),('Audit sai schema trong transaction','Notification trước đó rollback.')]
for i,(left,right) in enumerate(rows):
    y=2.05+i*1.03
    box(s,.7,y,4.3,.82,left,23,BLUE,True,LIGHT)
    box(s,5.32,y+.09,7.1,.75,right,23)
note(s,'EXPLAIN có TraversalNode và IndexNode; bằng chứng trên database kiểm thử riêng.')

s=slide('KẾT QUẢ KIỂM THỬ','Số liệu từ artifact đã có, không phải lượt chạy mới khi biên soạn PPT. recommendations-regression.xml: 167 tests, không fail/error/skip, 533,522 giây. recommendations-verified-playwright.xml: 20 tests, không fail/error/skip, 124,387 giây. arango-proof-tests.xml: 21 tests, không fail/error/skip, 141,334 giây. Các tập có thể chồng lặp, không cộng thành 188. Fixtures nghiệp vụ giả lập nhưng DB/HTTP thật; ArangoDB 3.12.12. Lượt browser đầu có 19 pass, 1 fail do email .test; sửa fixture sang .com rồi chạy lại đạt 20 pass.')
cards(s,[('167 / 167','Python integration\nArangoDB thật.'),('20 / 20','Browser end-to-end\nHTTP backend thật.'),('21 / 21','Graph / index / rollback\nBộ kiểm tra chuyên sâu.')])
note(s,'Artifact ngày 12/10/2026; các bộ có thể chồng lặp, không cộng số ca.')

s=slide('HIỆU NĂNG: KẾT QUẢ LOCAL','Nguồn: docs/reports/benchmark.json, ghi cho ngày 11/10/2026, trước cập nhật catalog public và giao diện gợi ý. Dataset synthetic có 1.134 users, 1.128 connections, 1.012 works_at; branching 10, depth 3. Mỗi query chạy 20 lần tuần tự, warm, local Docker. p95: job_matches 14,4849 ms; paths_depth_3 388,7014 ms; jobs_with_referral 580,4688 ms. Đây chưa là benchmark catalog hiện tại, không đo concurrency, SLA hoặc cluster. Cần chạy lại khi dữ liệu và cấu hình thay đổi.')
vals=[('Matching việc làm',14.485),('Đường kết nối · 3 hops',388.701),('Job + đường giới thiệu',580.469)]
for i,(label,value) in enumerate(vals):
    y=2.2+i*1.2
    box(s,.75,y,3.55,.72,label,23)
    width=value/600*6.9
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(4.45),Inches(y+.05),Inches(width),Inches(.55))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(BLUE);sh.line.fill.background()
    box(s,4.56+width,y+.01,2,.55,f'{value:.1f} ms'.replace('.',','),21,BLUE,True)
note(s,'p95 · 20 lần/query · synthetic · tuần tự · local; chưa là kiểm thử tải production.')

s=slide('DEMO SẢN PHẨM: 4 BƯỚC','Chuẩn bị Docker healthy, python -m database.cli verify. Demo dữ liệu public cho nguồn và matching; fixture riêng cho graph/referral. Chạy scripts/prove_arangodb.py để chứng minh mutation/index/rollback. Browser interactive: PowerShell env:JRN_E2E_REPORT_PREFIX=arango-interactive; .venv\\Scripts\\python.exe scripts/check_frontend.py --debug --grep "đăng ký →". Script tạo backend/Vite ports riêng và UUID DB/user, cleanup khi kết thúc. Không cần mutate main DB hoặc gán nhân viên giả cho public employers. Chi tiết docs/demo-guide.md, docs/arango-proof-guide.md.')
steps=[('01','Lưu hồ sơ + kỹ năng'),('02','Xem việc được gợi ý'),('03','Tìm đường A → B → C'),('04','Gửi → xử lý referral')]
for i,(num,label) in enumerate(steps):
    y=2.08+i*1.03
    box(s,.78,y,1.1,.8,num,32,'FFFFFF',True,BLUE,PP_ALIGN.CENTER)
    box(s,2.1,y+.08,10,.72,label,28)
note(s,'Nguồn public cho tin tuyển dụng; fixture riêng cho luồng graph và giới thiệu.')

s=slide('GIỚI HẠN VÀ HƯỚNG PHÁT TRIỂN','Nguồn: docs/demo-guide.md, public-data.md, deployment.md, security-summary.json và security-review-2026-10-11.md. Trọng số chưa được hiệu chỉnh bằng nhãn thực; chưa đánh giá Precision@K/NDCG. Catalog là snapshot, chưa có scheduler cập nhật; các công ty public chưa có nhân viên được xác minh. Upload CV, email và ATS chưa triển khai. Depth tối đa 3 và khóa transaction rộng; benchmark tuần tự. production_ready=false; image findings mức HIGH, TLS, tải và vận hành còn cần xử lý. Hướng phát triển là kế hoạch, chưa được tính là kết quả hoàn thành.')
cards(s,[('Gợi ý theo quy tắc','Thu thập phản hồi có consent\nvà hiệu chỉnh trọng số.'),('Dữ liệu hạn chế','Đồng bộ tin, xác minh\nemployment, thêm email/ATS.'),('MVP chạy local','Đo tải, tối ưu khóa\nvà hoàn thiện vận hành.')])
note(s,'Chưa đánh giá độ chính xác với người dùng thật; chưa nghiệm thu production.')

s=slide('KẾT LUẬN VÀ TRAO ĐỔI','ArangoDB lưu thực thể và quan hệ, thực thi AQL cho matching và traversal. Sản phẩm có luồng giao diện–API–database, điểm giải thích và workflow; bằng chứng kiểm thử graph/rollback trên local. Không suy ra xác suất tuyển dụng hoặc production readiness. Tài liệu chính thức đã đọc: https://docs.arango.ai/arangodb/3.12/graphs/ ; https://docs.arango.ai/arangodb/3.12/aql/ ; https://docs.arango.ai/arangodb/3.12/aql/graph-queries/traversals/ . Source: D:\\bigdata; docs/bao-cao-arangodb.md; docs/reports/; docs/arango-proof-guide.md. Thông tin nhóm và giảng viên còn cần bổ sung.')
bullets(s,['ArangoDB là lõi của matching và tìm đường graph.','Sản phẩm có giao diện, API và dữ liệu thật.','Có giải thích kết quả và bằng chứng kiểm thử.'],y=2.05,size=27,gap=1.05)
box(s,.8,5.32,11.8,.55,'CẢM ƠN THẦY CÔ VÀ CÁC BẠN',30,BLUE,True,align=PP_ALIGN.CENTER)
note(s,'Nguồn: tài liệu ArangoDB 3.12 • mã nguồn và reports dự án • báo cáo kèm theo.')

assert len(p.slides)==20
p.core_properties.title='Ứng dụng ArangoDB: gợi ý việc làm và đường kết nối giới thiệu'
p.core_properties.subject='Bài thuyết trình 20 slide theo mẫu nhận diện HUIT'
p.core_properties.author='[Cần bổ sung thông tin nhóm]'
for i,s in enumerate(p.slides):
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=p.slide_width and sh.top+sh.height<=p.slide_height,(i+1,sh.name)
    text=' '.join(sh.text for sh in s.shapes if sh.has_text_frame)
    assert 'Click to edit' not in text and 'Insert your content' not in text
    manifest[i]['visible_words']=len(text.split())
p.save(TARGET)
shutil.copy2(TARGET,OUT/'Thuyet_trinh_ArangoDB_HUIT.pptx')
(OUT/'slides-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'slides':len(p.slides),'target':str(TARGET),'backup':str(BACKUP),'notes':20,'max_visible_words':max(x['visible_words'] for x in manifest)},ensure_ascii=False))
