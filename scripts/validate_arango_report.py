"""Verify delivered Word/PDF structure and create a visual page index."""
from pathlib import Path
import json
import shutil
import xml.etree.ElementTree as ET

import pymupdf
from PIL import Image, ImageDraw
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/'docs/report-assets'
TARGET=Path(r'C:\Users\NGUYEN THANH NAM\Downloads\Mau bao cao_Do an_Khoa luan_Huong ung dung.docx')
d=Document(TARGET)
assert len(d.sections)==3
for sec in d.sections:
    for actual,expected in [(sec.page_width.cm,21),(sec.page_height.cm,29.7),
                            (sec.left_margin.cm,3.5),(sec.right_margin.cm,2.5),
                            (sec.top_margin.cm,2.5),(sec.bottom_margin.cm,2.5)]:
        assert abs(actual-expected)<0.01,(actual,expected)
assert d.styles['Normal'].font.name=='Times New Roman'
assert d.styles['Normal'].font.size.pt==13
assert abs(d.styles['Normal'].paragraph_format.line_spacing-1.3)<0.001
assert len(d.tables)==14
assert len(d.inline_shapes)==9  # Two template logos and seven report figures.
for tb in d.tables:
    assert tb.rows[0]._tr.find('.//'+qn('w:tblHeader')) is not None
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    assert r.font.name=='Times New Roman' and r.font.size.pt==12
footers=[sec.footer._element for sec in d.sections]
assert len(footers[0].findall('.//'+qn('w:instrText')))==0
assert [len(x.findall('.//'+qn('w:instrText'))) for x in footers[1:]]==[1,1]
assert d.sections[1]._sectPr.find(qn('w:pgNumType')).get(qn('w:fmt'))=='lowerRoman'
assert d.sections[2]._sectPr.find(qn('w:pgNumType')).get(qn('w:fmt'))=='decimal'
assert sum('TOC ' in (x.text or '') for x in d._element.findall('.//'+qn('w:instrText')))==3
pdf=pymupdf.open(ASSETS/'Bao_cao_ArangoDB.pdf')
for i,p in enumerate(pdf):
    assert len(p.get_text().strip())>35,('unexpected blank page',i+1)
    assert not any(x in p.get_text() for x in ['Error!','Bookmark not defined','Mục lục tự động'])
# Verify chapter content independently of generated TOC text.
heads=[p.text for p in d.paragraphs if p.style.name=='Heading 1']
assert sum(x.startswith('CHƯƠNG ') for x in heads)==5
assert 'MỞ ĐẦU' in heads and 'KẾT LUẬN' in heads and 'TÀI LIỆU THAM KHẢO' in heads
tests={}
for filename in ['recommendations-regression.xml','arango-proof-tests.xml','recommendations-verified-playwright.xml']:
    root=ET.parse(ROOT/'docs/reports'/filename).getroot()
    suites=root.findall('testsuite')
    attrs=suites[0].attrib if filename!='recommendations-verified-playwright.xml' else root.attrib
    tests[filename]={k:attrs[k] for k in ['tests','failures','errors','skipped']}
width,height=230,345
sheet=Image.new('RGB',(width*5,height*((len(pdf)+4)//5)),'#e5e5e5')
draw=ImageDraw.Draw(sheet)
for i,page in enumerate(pdf):
    pix=page.get_pixmap(matrix=pymupdf.Matrix(.37,.37))
    im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
    im.thumbnail((width-10,height-25))
    x=(i%5)*width+(width-im.width)//2;y=(i//5)*height
    sheet.paste(im,(x,y+20))
    draw.text((x,y+3),str(i+1),fill='black')
sheet.save(ASSETS/'all-pages-review.png')
for idx in [0,1,2,3,7,8,len(pdf)-3,len(pdf)-2,len(pdf)-1]:
    pdf[idx].get_pixmap(matrix=pymupdf.Matrix(1.2,1.2)).save(ASSETS/f'final-page-{idx+1:02}.png')
summary={'document':str(TARGET),'pages':len(pdf),'chapters':5,'tables':len(d.tables),
         'report_figures':7,'automatic_lists':3,'template_backup':str(TARGET.with_name(TARGET.stem+'_Mau_goc.docx')),
         'format_checks':'pass','field_checks':'pass','existing_test_evidence':tests,
         'new_system_tests_run_for_report':False,
         'administrative_fields_pending':['Loại báo cáo','Ngành','Sinh viên/MSSV/Lớp','Giảng viên hướng dẫn']}
(ASSETS/'report-validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
pdf.close()
shutil.copy2(TARGET,ASSETS/'Bao_cao_ArangoDB.docx')
pdf_target=TARGET.with_name('Bao_cao_ArangoDB_Goi_y_viec_lam.pdf')
if not pdf_target.exists():shutil.copy2(ASSETS/'Bao_cao_ArangoDB.pdf',pdf_target)
print(json.dumps(summary,ensure_ascii=False))
